"""XLSX result/history download regressions. Temporary DB, no network or mail."""
import io
import json
import sqlite3
import tempfile
import unittest
import zipfile
from pathlib import Path
from openpyxl import load_workbook
from agents.price_request_export import generate_xlsx, MIME
from agents.price_request_store import PriceRequestStore
from tests import test_price_request_routes as route_fixtures


def values(ws, row=2):
    return {cell.value: ws.cell(row, cell.column).value for cell in ws[1]}


def model(scope='a8'):
    return {'request': {'query': 'keytruda', 'scope': scope, 'id': 'a'*32, 'status': 'verified'},
            'snapshot_kind': 'verified', 'country_results': [{'country':'IT' if scope=='a8' else 'BE','status':'priced','row_ids':[1]}],
            'rows':[{'id':1,'country':'IT' if scope=='a8' else 'BE','product_name':'Keytruda','ingredient':'pembrolizumab',
                     'form_type':'injection','dosage_strength':'100 mg/4 ml','package_unit':'2 vials','pack_count':2,
                     'currency':'EUR','local_price':7596.68,'per_unit_local':3798.34,'exchange_rate':1500,
                     'factory_price_krw':5697510,'adjusted_price_krw':6811605.9309,
                     'source_url':'https://example.org/source?token=REDACT_ME&doc=12',
                     'source_label':'AIFA decision','raw_data': json.dumps({'source_type':'aifa_exfactory','aic':'044386035','evidence_pages':[6],
                         'price_tier_original':'Prezzo ex-factory (IVA esclusa)','vat_status':'excluded','observed_at':'2026-01-01T00:00:00Z'})}]}


class WorkbookTests(unittest.TestCase):
    def test_structure_values_provenance_and_safe_links(self):
        content, filename=generate_xlsx(model())
        self.assertTrue(filename.endswith('.xlsx'))
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            self.assertIsNone(z.testzip())
            self.assertFalse(any('vbaProject' in n or 'externalLinks/' in n for n in z.namelist()))
        wb=load_workbook(io.BytesIO(content),data_only=False)
        self.assertEqual(wb.sheetnames,['안내','국가별 결과','A8 가격 상세','출처와 증빙'])
        data=values(wb['A8 가격 상세'])
        self.assertEqual(data['원포장 가격 (현지통화)'],7596.68)
        self.assertEqual(data['포장 내 최소단위 수량'],2)
        self.assertEqual(data['최소단위 현지가'],3798.34)
        self.assertEqual(data['원가격 단계'],'공장도가')
        self.assertEqual(data['원가격 VAT'],'제외')
        evidence=values(wb['출처와 증빙'])
        self.assertEqual(evidence['AIC'],'044386035')
        self.assertEqual(evidence['PDF 페이지'],'[6]')
        self.assertNotIn('REDACT_ME',evidence['원문 URL'])
        self.assertIn('doc=12',evidence['원문 URL'])
        self.assertIsNotNone(wb['A8 가격 상세'].freeze_panes)
        self.assertIsNotNone(wb['A8 가격 상세'].auto_filter.ref)

    def test_source_text_cannot_become_formula(self):
        m=model()
        m['request']['query']='=HYPERLINK("bad")'
        m['rows'][0]['product_name']='=1+1'
        m['country_results'][0]['reason']='=WEBSERVICE("bad")'
        wb=load_workbook(io.BytesIO(generate_xlsx(m)[0]),data_only=False)
        self.assertEqual(wb['A8 가격 상세']['C2'].value,'=1+1')
        self.assertEqual(wb['A8 가격 상세']['C2'].data_type,'s')
        self.assertTrue(all(c.data_type!='f' for ws in wb for row in ws for c in row))

    def test_oecd_tiers_not_a8_adjusted(self):
        m=model('oecd')
        m['rows'][0].update(price_tier='ex_factory',ex_factory_price=7596.68,pharmacy_purchase_price=7600,grade='B',
                            derivation_note='Official regulation evidence',normalized_price=3798.34,total_pkg_mg=200,reference_mg=100)
        wb=load_workbook(io.BytesIO(generate_xlsx(m)[0]),data_only=False)
        data=values(wb['추가국가 가격 상세'])
        self.assertEqual(data['공장도가 VAT제외 (현지통화)'],7596.68)
        self.assertEqual(data['약국구입가 VAT제외 (현지통화)'],7600)
        self.assertNotIn('최소단위 A8 조정가 (원)',data)


class ExportRouteTests(unittest.TestCase):
    setUp=route_fixtures.RouteTests.setUp
    tearDown=route_fixtures.RouteTests.tearDown
    create=route_fixtures.RouteTests.create
    upload=route_fixtures.RouteTests.upload

    def workbook(self,response):
        self.assertEqual(response.status_code,200,response.json if response.is_json else '')
        self.assertEqual(response.mimetype,MIME)
        self.assertEqual(response.headers['Cache-Control'],'no-store')
        self.assertIn('attachment',response.headers['Content-Disposition'])
        return load_workbook(io.BytesIO(response.data),data_only=False)

    def test_request_download_owner_only(self):
        rid,_=self.upload()
        path=f'/api/price-requests/{rid}/export.xlsx'
        self.workbook(self.client.get(path,headers=self.user))
        self.assertEqual(self.client.get(path,headers=self.other).status_code,404)
        self.assertEqual(self.client.get(path,headers=self.worker).status_code,403)
        self.assertEqual(self.client.get(path).status_code,401)

    def test_snapshot_survives_deleted_cache(self):
        rid,_=self.upload()
        with sqlite3.connect(self.path) as c:
            c.execute('DELETE FROM foreign_drug_prices')
        wb=self.workbook(self.client.get(f'/api/price-requests/{rid}/export.xlsx',headers=self.user))
        self.assertEqual(wb['A8 가격 상세'].max_row,3)
        self.assertEqual(values(wb['A8 가격 상세'])['원포장 가격 (현지통화)'],7596.68)
        self.assertEqual(self.client.get('/api/price-history/export.xlsx?scope=a8&q=keytruda',headers=self.user).status_code,404)

    def test_verified_readback_freezes_values(self):
        rid,r=self.upload()
        read=self.client.get(f'/api/price-agent/requests/{rid}/readback',headers=self.headers).json
        b=dict(visible_row_ids=read['visible_row_ids'],ui_checked=True,evidence_ref='proof.png',dashboard_url=r['dashboard_url'])
        self.assertEqual(self.client.post(f'/api/price-agent/requests/{rid}/verify',headers=self.headers,json=b).status_code,200)
        with sqlite3.connect(self.path) as c:
            c.execute('UPDATE foreign_drug_prices SET local_price=999,adjusted_price_krw=999')
        wb=self.workbook(self.client.get(f'/api/price-requests/{rid}/export.xlsx',headers=self.user))
        self.assertEqual(values(wb['A8 가격 상세'])['원포장 가격 (현지통화)'],7596.68)
        summary={row[0].value:row[1].value for row in wb['안내'].iter_rows(min_row=2)}
        self.assertIn('화면 검증 당시',summary['자료 기준'])
        current=self.workbook(self.client.get('/api/price-history/export.xlsx?scope=a8&q=keytruda',headers=self.user))
        self.assertEqual(values(current['A8 가격 상세'])['원포장 가격 (현지통화)'],999)

    def test_pending_blocked_failure_only_downloadable(self):
        rid=self.create()
        path=f'/api/price-requests/{rid}/export.xlsx'
        self.assertEqual(self.client.get(path,headers=self.user).status_code,409)
        result=self.client.put(f'/api/price-agent/requests/{rid}/countries/IT',headers=self.headers,
                               json={'status':'failed','reason':'login unavailable','rows':[]})
        self.assertEqual(result.status_code,200)
        wb=self.workbook(self.client.get(path,headers=self.user))
        self.assertEqual(values(wb['국가별 결과'])['미확정·실패 사유'],'login unavailable')
        self.assertEqual(wb['A8 가격 상세'].max_row,1)

    def test_history_auth_empty_and_scope(self):
        path='/api/price-history/export.xlsx?scope=a8&q=keytruda'
        self.assertEqual(self.client.get(path,headers=self.user).status_code,404)
        self.assertEqual(self.client.get(path).status_code,401)
        self.assertEqual(self.client.get(path,headers=self.worker).status_code,403)
        self.assertEqual(self.client.get('/api/price-history/export.xlsx?scope=nope&q=x',headers=self.user).status_code,400)
        self.upload()
        self.workbook(self.client.get(path,headers=self.user))

    def test_oecd_history_and_request(self):
        store=PriceRequestStore(self.path)
        r=store.create('keytruda','oecd',['BE'],'user@example.org','user@example.org')
        lease=store.claim(r['id'],'worker')
        row=model('oecd')['rows'][0]
        row['observed_at']='2026-01-01T00:00:00Z'
        store.accept(r['id'],'BE','worker',lease['lease_token'],'priced',None,[row],'hash')
        for path in (f'/api/price-requests/{r["id"]}/export.xlsx','/api/price-history/export.xlsx?scope=oecd&q=keytruda'):
            wb=self.workbook(self.client.get(path,headers=self.user))
            self.assertEqual(values(wb['추가국가 가격 상세'])['원포장 가격 (현지통화)'],7596.68)


if __name__=='__main__':
    unittest.main()
