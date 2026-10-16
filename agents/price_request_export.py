"""Static XLSX exports of price-request snapshots or current saved history.
No scrape, FX refresh or formula recalculation occurs during export. Formula-like
source strings are written as text, never executable spreadsheet formulas.
"""
from __future__ import annotations
import io
import json
import math
import re
from datetime import datetime, timezone
from urllib.parse import urlsplit, urlunsplit, parse_qsl, urlencode

from agents.price_request_catalog import CATALOG

MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
LABELS = {x["code"]: x["label"] for scope in ("a8", "oecd") for x in CATALOG[scope]}
STATUS = {"pending": "대기", "priced": "가격 확인", "unconfirmed": "가격 미확정", "failed": "수집 실패",
          "queued": "접수 대기", "processing": "조사 중", "awaiting_verification": "화면 검증 대기",
          "verified": "화면 검증됨", "completed": "완료", "partial_completed": "부분 완료", "blocked": "차단"}
TIERS = {"ex_factory": "공장도가", "pharmacy_purchase": "약국구입가", "pharmacy_retail": "약국판매가", "reimbursement": "상환가", "other": "기타"}
VAT = {"included": "포함", "excluded": "제외", "unknown": "미확인"}


def _meta(row):
    raw = row.get("raw_data")
    if isinstance(raw, str):
        try:
            raw = json.loads(raw)
        except (ValueError, TypeError):
            raw = {}
    raw = raw if isinstance(raw, dict) else {}
    extra = row.get("extra") if isinstance(row.get("extra"), dict) else {}
    return {**raw, **extra, **{k: v for k, v in row.items() if v is not None}}


def _safe_url(value):
    if not isinstance(value, str):
        return None
    try:
        u = urlsplit(value)
        if u.scheme not in ("http", "https") or not u.hostname:
            return None
        blocked = re.compile(r"password|secret|token|api.?key|authorization|session|cookie", re.I)
        query = urlencode([(k, v) for k, v in parse_qsl(u.query, keep_blank_values=True) if not blocked.search(k)])
        host = u.hostname + (":" + str(u.port) if u.port else "")
        return urlunsplit((u.scheme, host, u.path, query, u.fragment))
    except ValueError:
        return None


def _scalar(v):
    if v is None:
        return None
    if isinstance(v, float) and not math.isfinite(v):
        return None
    if isinstance(v, (str, int, float, bool)):
        return v
    return json.dumps(v, ensure_ascii=False, default=str)


def _sheet(wb, name, columns, rows):
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    ws = wb.create_sheet(name)
    ws.append([title for key, title, width in columns])
    for values in rows:
        ws.append([_scalar(values.get(key)) for key, _, _ in columns])
    for row in ws:
        for cell in row:
            if isinstance(cell.value, str):
                # Protect against formula injection from product names, reasons or URLs.
                cell.value = ILLEGAL_CHARACTERS_RE.sub("", cell.value)[:32767]
                cell.data_type = "s"
            cell.font = Font(name="Calibri", size=11)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                cell.number_format = '#,##0.00####;[Red](#,##0.00####);"-"'
    for cell in ws[1]:
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="145C58")
    for index, (key, title, width) in enumerate(columns, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
        if key in ("id", "pack_count", "count"):
            for cell in list(ws.columns)[index - 1][1:]:
                cell.number_format = '#,##0'
        elif key in ("factory_ratio", "vat_rate", "distribution_margin"):
            for cell in list(ws.columns)[index - 1][1:]:
                cell.number_format = '0.0%'
        if key in ("source_url", "pdf_url"):
            for cell in list(ws.columns)[index - 1][1:]:
                url = _safe_url(cell.value)
                cell.value = url
                if url:
                    cell.hyperlink = url
                    cell.font = Font(name="Calibri", size=11, color="0563C1", underline="single")
    ws.freeze_panes = "E2" if len(columns) > 5 else "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 32
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:1"
    return ws


def generate_xlsx(data):
    """data={request,rows,country_results,snapshot_kind,...}; returns bytes+filename."""
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.creator = "MA AI Dossier"
    wb.properties.title = "해외약가 조사 결과"
    req = data["request"]
    scope = req["scope"]
    mode = {"verified": "화면 검증 당시 가격 스냅샷", "uploaded": "요청별 업로드 당시 결과 (화면 검증 전 또는 구형 기록)",
            "current_cache": "다운로드 시점의 현재 저장 캐시 (과거 요청 당시 결과가 아님)"}.get(data.get("snapshot_kind"), "요청별 결과")
    summary = [
        {"item": "제품/검색어", "value": req["query"]},
        {"item": "범위", "value": "A8 8개국" if scope == "a8" else "OECD·공단 추가국가 (A8 외 29개국 기준)"},
        {"item": "자료 기준", "value": mode},
        {"item": "요청 ID", "value": req.get("id") or "기존 검색 이력"},
        {"item": "요청 상태", "value": STATUS.get(req.get("status"), req.get("status"))},
        {"item": "요청일시", "value": req.get("created_at")},
        {"item": "대시보드 검증일시", "value": req.get("verified_at")},
        {"item": "가격 스냅샷 보관일시", "value": data.get("snapshot_captured_at")},
        {"item": "메일 발송 상태", "value": req.get("notification_status")},
        {"item": "엑셀 생성일시 (UTC)", "value": datetime.now(timezone.utc).isoformat(timespec="seconds")},
        {"item": "파일 성격", "value": "정적 결과표입니다. 표시된 값을 보존하며 수식·외부 데이터 연결·매크로를 포함하지 않습니다."},
        {"item": "비교 주의", "value": "원포장 가격·가격단계·VAT·제형·함량·포장을 구분하세요. A8과 추가국가를 합산하지 않습니다. 서로 다른 제형/함량의 평균을 계산하지 않습니다."},
        {"item": "날짜 주의", "value": "조회/관찰일·공시일·결정문일·가격 시행일은 다릅니다. 빈 칸은 미확인/미산출이지 0원이 아닙니다."},
        {"item": "국가 범위", "value": CATALOG["note"]},
    ]
    missing = data.get("missing_row_ids") or []
    if missing:
        summary.append({"item": "과거값 누락 경고", "value": "원행/보관본을 찾지 못한 ID: " + ", ".join(map(str, missing)) + ". 최신 가격으로 대체하지 않았습니다."})
    _sheet(wb, "안내", [("item", "항목", 28), ("value", "내용", 110)], summary)

    countries = data.get("country_results", [])
    country_rows = [{**x, "country_name": LABELS.get(x["country"], x["country"]), "status_label": STATUS.get(x.get("status"), x.get("status")),
                     "count": len(x.get("row_ids", []))} for x in countries]
    _sheet(wb, "국가별 결과", [("country", "국가코드", 12), ("country_name", "국가", 18), ("status_label", "조사 상태", 20),
                              ("count", "저장 가격행 수", 17), ("reason", "미확정·실패 사유", 70), ("updated_at", "상태 기록일시", 29)], country_rows)
    price_rows = []
    for row in data.get("rows", []):
        m = _meta(row)
        m["country_name"] = LABELS.get(m.get("country"), m.get("country"))
        tier = m.get("price_tier")
        m["tier_label"] = TIERS.get(tier, tier)
        m["vat_label"] = VAT.get(m.get("vat_status"), m.get("vat_status"))
        m["pages"] = m.get("pages") or m.get("evidence_pages")
        m["price_tier_original"] = m.get("price_tier_original") or m.get("price_kind")
        # Legacy A8 values without explicit native-tier metadata remain unknown.
        if scope == "a8" and not tier:
            source_type = m.get("source_type")
            if source_type == "aifa_exfactory" or m.get("price_kind") == "ex-factory":
                m["tier_label"] = "공장도가"
                m["price_tier_original"] = m.get("price_tier_original") or "Prezzo ex-factory"
                m["vat_label"] = m["vat_label"] or "제외"
            elif source_type == "redbook_wac":
                m["price_tier_original"] = "WAC"
        m["source_url"] = _safe_url(m.get("source_url"))
        m["pdf_url"] = _safe_url(m.get("pdf_url") or m.get("document_url"))
        if scope == "a8" and m.get("currency") == "JPY" and isinstance(m.get("exchange_rate"), (int, float)) and m["exchange_rate"] > 100:
            m["exchange_rate"] /= 100  # same per-one currency display convention as dashboard
        price_rows.append(m)
    common = [("id", "가격행 ID", 13), ("country_name", "국가", 17), ("product_name", "제품명", 28), ("ingredient", "성분명", 26),
              ("form_type", "제형 분류", 16), ("route", "투여경로/분류", 19), ("dosage_strength", "함량·규격 원문", 48),
              ("dosage_form", "제형 원문", 32), ("package_unit", "포장 원문", 28), ("pack_count", "포장 내 최소단위 수량", 23),
              ("local_price", "원포장 가격 (현지통화)", 26), ("currency", "통화", 12),
              ("tier_label", "원가격 단계", 20), ("price_tier_original", "가격 필드 원문", 34), ("vat_label", "원가격 VAT", 16)]
    if scope == "a8":
        extra = [("prezzo_pubblico", "IT 소비자가 참조 (EUR, 계산 미사용)", 37), ("per_unit_local", "최소단위 현지가", 22), ("unit_strength_mg", "최소단위 총함량 (mg)", 25),
                 ("exchange_rate", "적용환율 (원/현지통화 1)", 27), ("exchange_rate_from", "환율 기간 시작", 19), ("exchange_rate_to", "환율 기간 종료", 19),
                 ("factory_ratio", "공장도가 산출비율", 22), ("factory_price_krw", "최소단위 공장도가 (원)", 27),
                 ("vat_rate", "한국 A8 VAT 비율", 21), ("distribution_margin", "한국 A8 유통거래폭", 23),
                 ("adjusted_price_krw", "최소단위 A8 조정가 (원)", 28), ("adjusted_price_krw_normalized", "제형 내 보정 조정가 (원)", 29),
                 ("reference_strength_mg", "보정 기준함량 (mg)", 24), ("dose_norm_note", "제형·보정 주의", 42)]
    else:
        extra = [("ex_factory_price", "공장도가 VAT제외 (현지통화)", 31), ("pharmacy_purchase_price", "약국구입가 VAT제외 (현지통화)", 34),
                 ("grade", "근거 등급", 15), ("derivation_note", "단계 계산·출처 근거", 65), ("total_pkg_mg", "원포장 총함량 (mg)", 24),
                 ("reference_mg", "정규화 기준함량 (mg)", 27), ("normalized_price", "함량정규화 원가격 (현지통화)", 34)]
    _sheet(wb, "A8 가격 상세" if scope == "a8" else "추가국가 가격 상세", common + extra + [("observed_at", "실제 관찰일시", 30), ("effective_date", "가격 시행일", 20), ("searched_at", "DB 검색 세션 일시", 30)], price_rows)
    _sheet(wb, "출처와 증빙", [("id", "가격행 ID", 13), ("country_name", "국가", 17), ("product_name", "제품명", 27),
                            ("source_label", "출처명", 40), ("source_url", "원문 URL", 70), ("evidence_ref", "증빙 참조", 62),
                            ("pdf_url", "문서 URL", 70), ("pages", "PDF 페이지", 19), ("determina", "결정문 번호", 28),
                            ("document_date", "결정문 일자", 20), ("publication_date", "공시일", 20), ("effective_date", "가격 시행일", 20), ("aic", "AIC", 17)], price_rows)
    out = io.BytesIO()
    wb.save(out)
    slug = re.sub(r"[^A-Za-z0-9가-힣_-]+", "_", req["query"])[:70] or "prices"
    tag = req.get("id", "current_cache")
    return out.getvalue(), f"{slug}_{scope}_{tag}_prices.xlsx"
